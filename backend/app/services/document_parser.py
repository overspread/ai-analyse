import fitz
import docx
from pathlib import Path
from typing import List, Dict


def extract_text_from_pdf(file_path: str) -> List[Dict]:
    try:
        doc = fitz.open(file_path)
        if doc.page_count == 0:
            raise ValueError("PDF file has no pages")
        pages = []
        for page in doc:
            page_text = page.get_text().strip()
            if not page_text:
                continue
            pages.append({
                "page_number": page.number + 1,
                "text": page_text,
            })
        doc.close()
        if not pages:
            raise ValueError("No text could be extracted from PDF")
        return pages
    except Exception as e:
        raise ValueError(f"Failed to extract text from PDF: {str(e)}")


def extract_text_from_docx(file_path: str) -> List[Dict]:
    try:
        doc = docx.Document(file_path)
        if len(doc.paragraphs) == 0:
            raise ValueError("DOCX file has no paragraphs")
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        if not paragraphs:
            raise ValueError("No text could be extracted from DOCX")
        return [{
            "page_number": 1,
            "text": "\n".join(paragraphs).strip(),
        }]
    except Exception as e:
        raise ValueError(f"Failed to extract text from DOCX: {str(e)}")


def extract_text(file_path: str) -> List[Dict]:
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
